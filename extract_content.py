import sys
from docx import Document
import pdfplumber

def extract_docx(path):
    print(f"\n{'='*70}\nDOCX: {path}\n{'='*70}")
    try:
        doc = Document(path)
        # paragraphs
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                print(t)
        # tables
        for ti, table in enumerate(doc.tables):
            print(f"\n--- TABLE {ti+1} ---")
            for row in table.rows:
                cells = [c.text.strip().replace('\n',' ') for c in row.cells]
                print(' | '.join(cells))
    except Exception as e:
        print(f"ERROR reading docx {path}: {e}")

def extract_pdf(path, max_pages=40):
    print(f"\n{'='*70}\nPDF: {path}\n{'='*70}")
    try:
        with pdfplumber.open(path) as pdf:
            print(f"Total pages: {len(pdf.pages)}")
            for i, page in enumerate(pdf.pages):
                if i >= max_pages:
                    print(f"... (truncated at {max_pages} pages)")
                    break
                txt = page.extract_text() or ""
                txt = txt.strip()
                if txt:
                    print(f"\n--- PAGE {i+1} ---")
                    print(txt)
    except Exception as e:
        print(f"ERROR reading pdf {path}: {e}")

if __name__ == "__main__":
    files = sys.argv[1:]
    for f in files:
        if f.lower().endswith('.docx'):
            extract_docx(f)
        elif f.lower().endswith('.pdf'):
            extract_pdf(f)
        else:
            print(f"Unsupported: {f}")
